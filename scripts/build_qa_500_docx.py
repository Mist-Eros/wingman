"""Create the Word version of the 500-question Wingman prep guide."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
QA_DIR = ROOT / "docs" / "qa"
OUT = ROOT / "docs" / "Wingman-500-Intrebari-si-Raspunsuri.docx"

SECTIONS = (
    "Produs, problemă și pitch",
    "Roluri, conturi și permisiuni",
    "Taxonomie, meniu și raportare",
    "AI, Whisper și LLM",
    "Dispecerat, quorum și timp real",
    "Recompense, statistici și timere",
    "UI/UX, limbă și accesibilitate",
    "Backend, arhitectură și securitate",
    "Testare, demo și depanare",
    "Interviu, tradeoff-uri și roadmap",
)

FACTS = (
    "Stack: Flutter/Dart + FastAPI/Python + SQLite + REST/WSS + TLS 8443.",
    "AI local: Ollama qwen2.5:1.5b + faster-whisper small (CPU, int8).",
    "10 conturi demo: admin + 9 departamente; parole hash-uite salted SHA-256.",
    "16 tipuri de incidente; meniu 8×8×8 = 512 combinații.",
    "Urgențe: low / medium / high / critical (verde / galben / portocaliu / roșu).",
    "Quorum: un responder per departament primar. High/critical → pending admin → confirmare admin.",
    "Rewards: 2 puncte raportare, 5 puncte join + rezolvare; două categorii în UI.",
    "Seed demo: 1 low, 1 medium, 2 high, 1 critical; SF-FIRE cere 2 responderi.",
    "Teste: 198 Flutter, 24 backend, 111/111 smoke; flutter analyze curat.",
    "Build local mock: board în memorie, fără REST/WSS; switch Local/Live din login.",
)


def pairs_for(path: Path) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    q: str | None = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("Q: "):
            q = line[3:].strip()
        elif line.startswith("A: "):
            if q is None:
                raise ValueError(f"{path}: answer without question")
            pairs.append((q, line[3:].strip()))
            q = None
    if q is not None:
        raise ValueError(f"{path}: dangling question")
    return pairs


def main() -> int:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(10)

    title = doc.add_heading("Wingman — 500 de întrebări și răspunsuri", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Material de pregătire pentru demo și interviu — tehnic și non-tehnic"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)

    doc.add_paragraph(
        "Documentul conține exact 500 de întrebări cu răspunsuri scurte, "
        "structurate pe 10 secțiuni. Faptele reflectă build-ul demo curent "
        "(commit 55f52fc) și includ limitările pe care trebuie să le prezinți "
        "sincer."
    )

    doc.add_heading("Fapte-cheie de memorat", level=1)
    for fact in FACTS:
        doc.add_paragraph(fact, style="List Bullet")

    number = 0
    total = 0
    for section_index, (title, filename) in enumerate(
        zip(SECTIONS, [f"s{i:02d}.md" for i in range(1, 11)]),
        start=1,
    ):
        pairs = pairs_for(QA_DIR / filename)
        if len(pairs) != 50:
            raise SystemExit(f"{filename}: expected 50, found {len(pairs)}")
        doc.add_page_break()
        doc.add_heading(f"{section_index}. {title}", level=1)
        for question, answer in pairs:
            number += 1
            total += 1
            qp = doc.add_paragraph()
            run = qp.add_run(f"{number}. {question}")
            run.bold = True
            run.font.size = Pt(10.5)
            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Pt(14)
            label = ap.add_run("Răspuns: ")
            label.bold = True
            ap.add_run(answer)

    if total != 500:
        raise SystemExit(f"expected 500, wrote {total}")
    doc.save(OUT)
    print(f"wrote {OUT} with {number} Q/A pairs")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
